import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import io
import re
import os
import unicodedata
import pdfplumber
from datetime import datetime

APP_VERSION = "v2.3"

# ==========================================
# 1. 基本設定
# ==========================================
st.set_page_config(
    page_title="雪茄批發報價自動整理系統",
    page_icon="🍂",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ==========================================
# 全域 CSS 樣式
# ==========================================
st.markdown("""
<style>
    .block-container { padding-top: 1rem; }

    /* ===== Hero 橫幅 ===== */
    .hero {
        background: linear-gradient(135deg, #0F2027 0%, #203A43 40%, #2C5364 100%);
        padding: 1.8rem 2.5rem;
        border-radius: 16px;
        margin-bottom: 1.4rem;
        position: relative;
        overflow: hidden;
    }
    .hero::before {
        content: '';
        position: absolute;
        top: -50%; right: -15%;
        width: 400px; height: 400px;
        background: radial-gradient(circle, rgba(255,255,255,0.06) 0%, transparent 70%);
        border-radius: 50%;
    }
    .hero h1 { color: #FFFFFF; font-size: 2.2rem; font-weight: 700; margin: 0 0 0.4rem 0; }
    .hero .subtitle { color: rgba(255,255,255,0.85); font-size: 1.1rem; margin: 0; line-height: 1.6; }
    .hero .badge {
        display: inline-block;
        background: rgba(255,255,255,0.15);
        border: 1px solid rgba(255,255,255,0.25);
        color: #FFFFFF;
        padding: 0.2rem 0.8rem;
        border-radius: 20px;
        font-size: 0.9rem;
        margin-top: 0.8rem;
    }

    /* ===== KPI 卡片 ===== */
    [data-testid="stMetric"] {
        background: #FFFFFF;
        border: 1px solid #E3E8F0;
        border-left: 5px solid #2C5364;
        border-radius: 12px;
        padding: 1rem 1.2rem;
        box-shadow: 0 2px 8px rgba(0,0,0,0.04);
        transition: transform 0.2s, box-shadow 0.2s;
    }
    [data-testid="stMetric"]:hover { transform: translateY(-2px); box-shadow: 0 4px 16px rgba(0,0,0,0.08); }
    [data-testid="stMetricLabel"] { color: #6B7280 !important; font-size: 1rem !important; font-weight: 600 !important; }
    [data-testid="stMetricValue"] { font-size: 1.7rem !important; font-weight: 700 !important; color: #1F2937 !important; }

    /* ===== 側邊欄 ===== */
    section[data-testid="stSidebar"] > div:first-child {
        background: linear-gradient(180deg, #EEF2F6 0%, #FFFFFF 100%);
    }
    .sidebar-title {
        background: linear-gradient(135deg, #0F2027, #2C5364);
        color: white;
        padding: 0.8rem 1rem;
        border-radius: 10px;
        font-size: 1.15rem;
        font-weight: 600;
        margin-bottom: 0.8rem;
        text-align: center;
    }
    .step-label {
        font-size: 0.95rem;
        font-weight: 700;
        color: #2C5364;
        letter-spacing: 1px;
        margin: 0.8rem 0 0.3rem 0;
    }

    /* ===== 下載按鈕 ===== */
    .stDownloadButton > button {
        background: linear-gradient(135deg, #1F4E78, #2C5E3B) !important;
        color: white !important;
        border: none !important;
        border-radius: 10px !important;
        font-weight: 600 !important;
        font-size: 1.05rem !important;
        padding: 0.5rem 1rem !important;
        box-shadow: 0 2px 8px rgba(31,78,120,0.25) !important;
        transition: all 0.25s !important;
    }
    .stDownloadButton > button:hover { transform: translateY(-1px) !important; box-shadow: 0 4px 14px rgba(31,78,120,0.4) !important; }

    /* ===== 分區標題 ===== */
    .section-header {
        display: flex; align-items: center; gap: 0.6rem;
        padding: 0.6rem 0; margin: 1rem 0 0.6rem 0;
        border-bottom: 2px solid #E5E7EB;
    }
    .section-header .icon { font-size: 1.6rem; }
    .section-header .text { font-size: 1.3rem; font-weight: 700; color: #1F2937; }
    .section-header .desc { font-size: 1rem; color: #6B7280; margin-left: auto; }

    /* ===== 功能卡片 (空狀態) ===== */
    .empty-state { text-align: center; padding: 3rem 2rem 2rem 2rem; color: #9CA3AF; }
    .empty-state .icon { font-size: 4rem; margin-bottom: 0.8rem; }
    .empty-state .title { font-size: 1.5rem; font-weight: 600; color: #4B5563; }
    .empty-state .desc { font-size: 1.1rem; margin-top: 0.5rem; }
    .feature-card {
        background: #FFFFFF;
        border: 1px solid #E5E7EB;
        border-radius: 14px;
        padding: 1.5rem 1.2rem;
        height: 100%;
        box-shadow: 0 2px 8px rgba(0,0,0,0.03);
    }
    .feature-card .f-icon { font-size: 2.2rem; }
    .feature-card .f-title { font-size: 1.15rem; font-weight: 700; color: #1F2937; margin: 0.5rem 0; }
    .feature-card .f-desc { font-size: 1rem; color: #6B7280; line-height: 1.6; }

    /* ===== 狀態列 ===== */
    .status-bar {
        background: linear-gradient(90deg, #ECFDF5, #F0FDF4);
        border: 1px solid #A7F3D0;
        border-radius: 10px;
        padding: 0.8rem 1.2rem;
        display: flex; align-items: center; gap: 0.8rem; flex-wrap: wrap;
        margin: 0.6rem 0 0.8rem 0;
    }
    .status-bar .dot { width: 10px; height: 10px; background: #10B981; border-radius: 50%; animation: pulse 2s infinite; }
    @keyframes pulse { 0%, 100% { opacity: 1; } 50% { opacity: 0.4; } }
    .status-bar .text { font-size: 1.05rem; color: #065F46; font-weight: 600; }
    .chip {
        display: inline-block;
        background: #FFFFFF;
        border: 1px solid #A7F3D0;
        color: #065F46;
        border-radius: 16px;
        padding: 0.15rem 0.8rem;
        font-size: 0.95rem;
        font-weight: 500;
    }

    /* ===== Tabs ===== */
    .stTabs [data-baseweb="tab-list"] { gap: 0; background: #F3F4F6; border-radius: 12px; padding: 4px; }
    .stTabs [data-baseweb="tab"] { border-radius: 8px; padding: 0.6rem 1.2rem; font-weight: 600; font-size: 1.05rem; }
    .stTabs [aria-selected="true"] { background: white !important; box-shadow: 0 1px 4px rgba(0,0,0,0.1); }

    /* ===== 放大 Dataframe 內容 ===== */
    [data-testid="stDataFrame"] { font-size: 1.05rem; }
    
    /* ===== 頁腳 ===== */
    .footer {
        text-align: center; padding: 1.5rem 0; color: #9CA3AF; font-size: 0.95rem;
        border-top: 1px solid #F3F4F6; margin-top: 2rem;
    }
</style>
""", unsafe_allow_html=True)

# ==========================================
# 2. 品牌與供應商設定 (動態載入 brand_rules.csv)
# ==========================================
RULES_FILE = 'brand_rules.csv'

DEFAULT_BRAND_RULES = [
    # 順序很重要：例如 "好友" 必須在 "蒙特" 前面，避免 "好友蒙特利" 誤判為蒙特
    {"品牌分類": "富恩特", "關鍵字": "富恩特, fuente, 海明威"},
    {"品牌分類": "我的父親", "關鍵字": "我的父亲, 我的父親, my father"},
    {"品牌分類": "多明尼加之花", "關鍵字": "多米尼加之花, 多明尼加之花, flor dominicana"},
    {"品牌分類": "獅子王", "關鍵字": "狮子王, 獅子王, 拉奥罗拉, aurora"},
    {"品牌分類": "好友", "關鍵字": "好友, hoyo, epicure, 逍遥, 赛科, 聖胡安, 帕尔马斯"},
    {"品牌分類": "潑辣", "關鍵字": "泼辣, 潑辣, 波尔, larranaga"},
    {"品牌分類": "高希霸", "關鍵字": "高希霸, cohiba, bhk, 世纪, 半世纪, 魔术师, 导师"},
    {"品牌分類": "千里達", "關鍵字": "千里达, 千里達, 特立尼达, trinidad, 3t, 雷耶斯, 暗礁"},
    {"品牌分類": "蒙特", "關鍵字": "蒙特, montecristo, 艾蒙多, 1935, 大仲马, 蒙特2号, 蒙特4号"},
    {"品牌分類": "帕特加斯", "關鍵字": "帕特加斯, partagas, 路西塔尼亚, d4, p2, e2"},
    {"品牌分類": "羅密歐與茱麗葉", "關鍵字": "罗密欧, 羅密歐, romeo, 短丘, 宽丘, 丘比特"},
    {"品牌分類": "烏普曼", "關鍵字": "乌普曼, 烏普曼, upmann, 玛瑙, 鉴赏家, 半皇冠"},
    {"品牌分類": "玻利瓦", "關鍵字": "玻利瓦, bolivar"},
    {"品牌分類": "雷蒙阿隆尼", "關鍵字": "雷蒙, allones"},
    {"品牌分類": "多爾賽碼頭", "關鍵字": "多尔塞, quai, d'orsay, 码头"},
    {"品牌分類": "潘趣", "關鍵字": "潘趣, punch"},
    {"品牌分類": "庫阿巴", "關鍵字": "库阿巴, cuaba"},
    {"品牌分類": "威古洛", "關鍵字": "威古洛, vegueros"},
    {"品牌分類": "胡安洛佩斯", "關鍵字": "胡安, juan lopez"},
    {"品牌分類": "外交官", "關鍵字": "外交官, diplomaticos"},
    {"品牌分類": "比雅達", "關鍵字": "比亚达, 比雅達, piedra, 猎人"},
    {"品牌分類": "金特羅", "關鍵字": "金特罗, quintero, 挚爱"},
    {"品牌分類": "世界之王", "關鍵字": "世界之王, rey del mundo"},
    {"品牌分類": "卡諾之花", "關鍵字": "卡诺之花, flor de cano"},
    {"品牌分類": "古巴榮耀", "關鍵字": "古巴荣耀, gloria cubana"},
    {"品牌分類": "拉斐爾", "關鍵字": "拉斐尔, rafael gonzalez"},
    {"品牌分類": "豐塞卡", "關鍵字": "丰塞卡, fonseca"},
    {"品牌分類": "羅賓納", "關鍵字": "罗宾纳, robaina"},
    {"品牌分類": "大衛杜夫", "關鍵字": "大卫杜夫, davidoff"},
    {"品牌分類": "潮牌CAO", "關鍵字": "潮牌, cao"},
]

def load_brand_rules():
    if not os.path.exists(RULES_FILE):
        pd.DataFrame(DEFAULT_BRAND_RULES).to_csv(RULES_FILE, index=False, encoding='utf-8-sig')
    try:
        return pd.read_csv(RULES_FILE)
    except Exception:
        return pd.DataFrame(DEFAULT_BRAND_RULES)

df_brand_rules = load_brand_rules()
TARGET_BRANDS = df_brand_rules['品牌分類'].tolist()

VENDOR_RULES = {
    "天空": "古巴天空",
    "千源": "千源古巴行",
    "古巴之家": "古巴之家",
    "818": "818同行批發",
    "壹茄": "壹茄批發",
    "寰宇": "寰宇之家",
    "维嘉": "維嘉批發",
    "維嘉": "維嘉批發",
    "行货": "行貨雪茄報價",
    "行貨": "行貨雪茄報價",
    "2026": "澳門批發庫存表",
    "澳门批发": "澳門批發庫存表",
    "澳門批發": "澳門批發庫存表",
    "庫存": "庫存報價(09.29)",
    "库存": "庫存報價(09.29)",
}

def get_standard_vendor(filename):
    """根據檔名關鍵字自動標準化供應商名稱"""
    for key, standard_name in VENDOR_RULES.items():
        if key in filename:
            return standard_name
    name = filename.rsplit('.', 1)[0]
    name = re.sub(r'\d{1,2}\.\d{1,2}', '', name)
    name = re.sub(r'20\d{2}', '', name)
    for noise in ['报价', '报價', '港币', '港幣', '最新', '澳门', '澳門', '批发', '批發', '价格', '價格', '号', '號', '(1)', '（1）']:
        name = name.replace(noise, '')
    return name.strip(' -_') or filename.rsplit('.', 1)[0]

def classify_brand(pname):
    """根據 CSV 設定的品牌關鍵字，嚴格匹配 (順序依 CSV 排列)"""
    text = str(pname).lower()

    # 針對「好友蒙特利」的特例處理（避免誤判為蒙特）
    is_hoyo = any(k in text for k in ['好友', 'hoyo', 'epicure', '逍遥', '赛科', '聖胡安', '帕尔马斯'])
    is_monte = any(k in text for k in ['蒙特克里斯托', 'montecristo', '蒙特2号', '蒙特4号'])
    if is_hoyo and not is_monte:
        return '好友'

    for _, row in df_brand_rules.iterrows():
        brand = row['品牌分類']
        if brand == '好友': continue # 特例已處理
        
        # 解析關鍵字清單，忽略空值
        keywords = [k.strip().lower() for k in str(row['關鍵字']).split(',') if k.strip()]
        if any(k in text for k in keywords):
            return brand

    return '其他品牌'

# ==========================================
# 3. 解析引擎與支數擷取
# ==========================================
CHANNEL_CLEAN_MAP = {
    "国营": "古巴國營", "古巴国营": "古巴國營", "欧水": "歐洲水貨", "西行": "西行/西班牙", "西班牙": "西行/西班牙",
    "英国免税": "英國免稅", "英免": "英國免稅", "英国": "英國免稅", "瑞士": "瑞士", "瑞行": "瑞士",
    "葡萄牙": "葡萄牙", "葡萄牙行": "葡萄牙", "德国": "德國", "德行": "德國", "加勒比": "加勒比",
    "加勒比海": "加勒比", "安道尔": "安道爾", "腓尼基": "腓尼基", "墨西哥": "墨西哥", "荷兰": "荷蘭",
    "爱尔兰": "愛爾蘭", "加拿大": "加拿大", "意大利": "意大利", "法行": "法行", "巴西": "巴西",
    "国营，欧水": "國營/歐水", "国营/欧水": "國營/歐水", "欧水，国营": "國營/歐水", "欧水/国营": "國營/歐水",
    "国营，西行": "國營/西行", "国营/西行": "國營/西行", "国营，腓尼基": "國營/腓尼基", "国营/腓尼基": "國營/腓尼基",
    "欧水，西行": "歐水/西行", "欧水/西行": "歐水/西行", "腓尼基/加勒比": "腓尼基/加勒比", "国营，英免": "國營/英免",
    "西行3欧水7": "混合調貨(西行+歐水)", "西行，腓尼基": "西行/腓尼基", "瑞士1欧水38": "混合調貨(瑞士+歐水)",
    "机制雪茄": "機制雪茄", "16年": "2016年份", "20年": "2020年份", "21年": "2021年份",
    "22年": "2022年份", "23年": "2023年份", "西行/22年": "西行/2022年份", "西行/24年": "西行/2024年份", "英飞烽": "英飛烽"
}

def clean_channel_name(val):
    s = str(val).strip() if val is not None else ""
    # 嚴格過濾：如果不在提供的字典裡，一律視為「未標註」
    return CHANNEL_CLEAN_MAP.get(s, "未標註")
def extract_quantity(name):
    """從品名規格中精準提取單盒支數"""
    s = str(name).strip()

    # 1. 乘法規格：如 10'S X 10, 5x5, 6*10, 5*4支
    m_mult = re.search(r'(\d+)\s*[\'’]?[sS]?\s*[xX*×]\s*(\d+)', s)
    if m_mult:
        return int(m_mult.group(1)) * int(m_mult.group(2))

    # 2. 中文支數：如 25支、10 支、25裝
    m_chi = re.search(r'(\d+)\s*(支|裝|装|/盒|盒)', s, re.IGNORECASE)
    if m_chi:
        return int(m_chi.group(1))

    # 3. 英文 S 規格：如 25S, 10S, 50s, 25'S
    m_s = re.search(r'(\d+)\s*[\'’]?[sS]\b', s)
    if m_s:
        val = int(m_s.group(1))
        if val in [1, 2, 3, 4, 5, 6, 8, 10, 12, 14, 15, 16, 18, 20, 24, 25, 30, 40, 50, 60, 66, 88, 100]:
            return val

    # 4. 優先匹配標準括號形式 (含 OCR 錯字容錯)：如 (25)、(25支)、( 15 )、(12枚)、(20页)
    match = re.search(r'[（\(]\s*(\d{1,2})\s*(?:支|枚|页|分)?\s*[）\)]', s)
    if match:
        return int(match.group(1))

    # 5. 處理左括號被 OCR 辨識成數字 1 的情況：如 "115 )"、"110 )"
    match = re.search(r'(?:^|[^\d])1(\d{2})\s*[）\)]', s)
    if match:
        return int(match.group(1))

    # 6. 處理未閉合括號：如 "(50 英国免"、"(16 水古"
    match = re.search(r'[（\(]\s*(\d{1,2})\s*(?:[^\d\n]|$)', s)
    if match:
        # 排除年份如 "(24年"
        if not re.search(r'[（\(]\s*\d{1,2}\s*年', s):
            return int(match.group(1))

    # 5. 特殊術語慣例
    if '單支' in s or '单支' in s: return 1
    if '宽短' in s or '寬短' in s: return 60
    if '短号' in s or '短號' in s: return 100

    return None
PRICE_CLEAN_RE = re.compile(r'[,$\s]|HKD|RMB|USD|EUR|¥|￥', flags=re.IGNORECASE)

def is_price(val):
    """判斷一個字串是否為有效的價格數值"""
    try:
        cv = PRICE_CLEAN_RE.sub('', str(val))
        return cv.replace('.', '', 1).isdigit() and float(cv) > 50
    except Exception:
        return False

def extract_price(val):
    """從字串中萃取價格數值"""
    return int(float(PRICE_CLEAN_RE.sub('', str(val))))

def _split_chunk(chunk, out_rows):
    """若同一段資料內有多個價格（並排的報價單），自動切成多筆"""
    p_count = sum(1 for v in chunk if is_price(v))
    if p_count > 1:
        size = max(1, len(chunk) // p_count)
        for i in range(p_count):
            out_rows.append(chunk[i*size:(i+1)*size])
    elif len(chunk) >= 2:
        out_rows.append(chunk)

def try_structured_parse(df):
    """嘗試根據標題列來進行結構化解析，以精確抓取「數量」與「價格」"""
    header_row_idx = -1
    for i in range(min(15, len(df))):
        row_str = "".join([str(x).lower() for x in df.iloc[i] if pd.notna(x)])
        if ("名" in row_str or "品" in row_str) and ("价" in row_str or "價" in row_str or "港" in row_str):
            header_row_idx = i
            break

    if header_row_idx == -1: return None

    headers = [str(x).lower().replace('\n', '').strip() if pd.notna(x) else "" for x in df.iloc[header_row_idx]]
    
    name_indices = []
    price_idx, stock_idx = -1, -1
    
    for j, h in enumerate(headers):
        if any(kw in h for kw in ["名", "品名", "规格", "英文", "中文"]): name_indices.append(j)
        if any(kw in h for kw in ["价", "價", "港币", "hkd", "rmb"]): price_idx = j
        if any(kw in h for kw in ["量", "庫存", "现货", "數量", "qty"]): stock_idx = j
            
    if not name_indices or price_idx == -1: return None
        
    structured_records = []
    for i in range(header_row_idx + 1, len(df)):
        row = df.iloc[i]
        name_parts = [unicodedata.normalize('NFKC', str(row[j]).strip()) for j in name_indices if pd.notna(row[j]) and str(row[j]).strip()]
        name = " ".join(name_parts)
        price_val = row[price_idx]
        
        if pd.notna(price_val) and is_price(price_val) and name:
            stock_qty = None
            if stock_idx != -1 and pd.notna(row[stock_idx]):
                sv_str = str(row[stock_idx]).replace('支', '').strip()
                if sv_str.isdigit(): stock_qty = int(sv_str)

            structured_records.append({
                "name": name,
                "price": extract_price(price_val),
                "stock_qty": stock_qty,
                "raw_row": [unicodedata.normalize('NFKC', str(x).strip()) for x in row if pd.notna(x) and str(x).strip()]
            })
            
    return structured_records if structured_records else None

@st.cache_data(show_spinner=False)
def parse_file_v5(file_bytes, file_ext):
    """解析單一檔案，回傳不含供應商的記錄清單（依檔案內容快取，重複上傳不需重算）"""
    structured_data = []
    raw_rows = []
    
    if file_ext in ['xlsx', 'xls']:
        sheets = pd.read_excel(io.BytesIO(file_bytes), header=None, sheet_name=None)
        for df_raw in sheets.values():
            s_records = try_structured_parse(df_raw)
            if s_records:
                structured_data.extend(s_records)
            else:
                for _, row in df_raw.iterrows():
                    current_chunk = []
                    for x in row:
                        val = str(x).strip()
                        if pd.notna(x) and val != "" and val.lower() != 'nan':
                            current_chunk.append(unicodedata.normalize('NFKC', val))
                        elif current_chunk:
                            _split_chunk(current_chunk, raw_rows)
                            current_chunk = []
                    if current_chunk:
                        _split_chunk(current_chunk, raw_rows)
    elif file_ext == 'pdf':
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables()
                if tables:
                    for table in tables:
                        for row in table:
                            if row:
                                cells = [unicodedata.normalize('NFKC', str(x).strip()) for x in row if x and str(x).strip()]
                                if cells:
                                    _split_chunk(cells, raw_rows)
                else:
                    text = page.extract_text()
                    if text:
                        for line in text.split('\n'):
                            cols = [unicodedata.normalize('NFKC', x.strip()) for x in re.split(r'\s{2,}|\t', line) if x.strip()]
                            if len(cols) == 1:
                                raw_rows.append(cols)
                            elif cols:
                                _split_chunk(cols, raw_rows)

    # 將非結構化抓取的 raw_rows 轉換為類似 structured_data 的格式
    for vals in raw_rows:
        vals = list(vals)
        if len(vals) == 1:
            vals = [x.strip() for x in re.split(r'\s{2,}|\t', vals[0]) if x.strip()]
        if len(vals) < 2:
            continue

        price, name = None, None
        for v in reversed(vals):
            if is_price(v):
                price = extract_price(v)
                vals.remove(v)
                break

        if price:
            candidates = [v for v in vals if not v.isdigit()]
            if candidates:
                name = max(candidates, key=len)
                vals.remove(name)

        if name and price:
            structured_data.append({
                "name": name,
                "price": price,
                "stock_qty": None,
                "raw_row": vals # 剩下的元素用來找渠道
            })

    records = []
    for r in structured_data:
        name = r['name']
        price = r['price']
        vals = r['raw_row']
        
        qty = extract_quantity(name + " " + " ".join(vals))
        brand = classify_brand(name + " " + " ".join(vals))
        
        origin_candidates = [v for v in vals if len(v) < 15 and not v.isdigit()
                             and not any(b in v for b in TARGET_BRANDS)]
        original_origin = " ".join(origin_candidates) if origin_candidates else ""
        
        final_channel = clean_channel_name(original_origin)
        
        if final_channel != "未標註" and original_origin in name:
            name = name.replace(original_origin, '').strip()
            
        records.append({
            "品牌分類": brand, 
            "品名規格": name, 
            "規格支數": qty,
            "庫存盒數": r.get('stock_qty'),
            "價格": price, 
            "渠道/產地": final_channel
        })
    return records

def build_best_price(df):
    """建立「最低價排行」：同品名在 2 家以上供應商有報價時，計算最低價與價差"""
    if df.empty:
        return pd.DataFrame()
    d = df.reset_index(drop=True)
    keys = ['品牌分類', '品名規格']
    agg = d.groupby(keys).agg(報價家數=('供應商', 'nunique'), 最低價=('價格', 'min'), 最高價=('價格', 'max'))
    best_vendor = d.loc[d.groupby(keys)['價格'].idxmin()].set_index(keys)['供應商'].rename('最低價供應商')
    out = agg.join(best_vendor).reset_index()
    out = out[out['報價家數'] >= 2].copy()
    out['價差'] = out['最高價'] - out['最低價']
    out['價差%'] = (out['價差'] / out['最低價'] * 100).round(1)
    out = out[['品牌分類', '品名規格', '最低價供應商', '最低價', '最高價', '價差', '價差%', '報價家數']]
    return out.sort_values('價差', ascending=False).reset_index(drop=True)

def build_pivot(df):
    if df.empty:
        return pd.DataFrame()
    return df.pivot_table(index=['品牌分類', '品名規格'], columns='供應商',
                          values='價格', aggfunc='min').astype('Int64')

# ==========================================
# 4. Excel 匯出（多工作表）
# ==========================================
FONT_NAME = 'Microsoft JhengHei'

def _autofit(ws):
    for col in ws.columns:
        max_len = max((sum(2 if ord(ch) > 127 else 1 for ch in str(c.value or '')) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 3, 10), 48)

def _style_header(ws, ncols, fill):
    header_font = Font(name=FONT_NAME, size=11, bold=True, color='FFFFFF')
    for col_num in range(1, ncols + 1):
        c = ws.cell(row=1, column=col_num)
        c.font, c.fill = header_font, fill
        c.alignment = Alignment(horizontal='center', vertical='center')

@st.cache_data(show_spinner=False)
def generate_excel(df_list, df_best, pivot):
    """產出含 4 個工作表的 Excel：全品項清單、跨店比價矩陣、最低價排行、品牌統計"""
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    body_font = Font(name=FONT_NAME, size=10)
    bold_font = Font(name=FONT_NAME, size=10, bold=True)
    navy_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    green_fill = PatternFill(start_color='2C5E3B', end_color='2C5E3B', fill_type='solid')
    best_fill = PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid')
    thin_border = Border(*(Side(style='thin', color='D9D9D9'),) * 4)

    # --- 工作表 1：全品項清單（正式表格，可直接插入交叉分析篩選器） ---
    ws_list = wb.create_sheet(title='全品項清單')
    ws_list.append(['供應商', '品牌分類', '品名規格', '規格支數', '庫存現貨', '整盒批發價', '渠道/產地'])
    for _, row in df_list.iterrows():
        qty = row.get('規格支數')
        qty = int(qty) if pd.notna(qty) else ''
        stock = row.get('庫存盒數')
        stock = int(stock) if pd.notna(stock) else ''
        ws_list.append([row['供應商'], row['品牌分類'], row['品名規格'], qty, stock, int(row['價格']), row.get('渠道/產地', '')])
    for c in range(1, 8):
        ws_list.cell(row=1, column=c).font = Font(name=FONT_NAME, size=11, bold=True, color='FFFFFF')
        ws_list.cell(row=1, column=c).alignment = Alignment(horizontal='center', vertical='center')
    for r in range(2, ws_list.max_row + 1):
        for c in range(1, 8):
            cell = ws_list.cell(row=r, column=c)
            cell.font = bold_font if c == 2 else body_font
            if c in [1, 2, 4, 5, 7]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c == 6:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '#,##0'
    if ws_list.max_row > 1:
        tab = Table(displayName="DataList", ref=f"A1:G{ws_list.max_row}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        ws_list.add_table(tab)
    ws_list.freeze_panes = 'A2'
    _autofit(ws_list)

    # --- 工作表 2：跨店比價矩陣（最低價綠底標示） ---
    ws_pv = wb.create_sheet(title='跨店比價矩陣')
    if not pivot.empty:
        vendors = list(pivot.columns)
        ws_pv.append(['品牌分類', '品名規格'] + vendors)
        for (brand, pname), row in pivot.iterrows():
            values = [None if pd.isna(v) else int(v) for v in row]
            ws_pv.append([brand, pname] + values)
            valid = [v for v in values if v is not None]
            min_v = min(valid) if valid else None
            r = ws_pv.max_row
            for i, v in enumerate(values):
                cell = ws_pv.cell(row=r, column=3 + i)
                cell.number_format = '#,##0'
                cell.border = thin_border
                if v is not None and v == min_v and len(valid) > 1:
                    cell.fill, cell.font = best_fill, bold_font
                else:
                    cell.font = body_font
            ws_pv.cell(row=r, column=1).font = bold_font
            ws_pv.cell(row=r, column=2).font = body_font
        _style_header(ws_pv, 2 + len(vendors), navy_fill)
        ws_pv.auto_filter.ref = f"A1:{get_column_letter(2 + len(vendors))}{ws_pv.max_row}"
        ws_pv.freeze_panes = 'C2'
        _autofit(ws_pv)

    # --- 工作表 3：最低價排行 ---
    ws_best = wb.create_sheet(title='最低價排行')
    if not df_best.empty:
        ws_best.append(list(df_best.columns))
        for row in df_best.itertuples(index=False):
            ws_best.append(list(row))
        _style_header(ws_best, len(df_best.columns), green_fill)
        for r in range(2, ws_best.max_row + 1):
            for c in range(1, len(df_best.columns) + 1):
                cell = ws_best.cell(row=r, column=c)
                cell.font, cell.border = body_font, thin_border
                if c in [4, 5, 6]:
                    cell.number_format = '#,##0'
                elif c == 7:
                    cell.number_format = '0.0"%"'
            ws_best.cell(row=r, column=3).font = bold_font
        ws_best.auto_filter.ref = f"A1:{get_column_letter(len(df_best.columns))}{ws_best.max_row}"
        ws_best.freeze_panes = 'A2'
        _autofit(ws_best)
    else:
        ws_best.append(['目前沒有同時出現在 2 家以上供應商的品項'])

    # --- 工作表 4：品牌分類統計 ---
    ws_sum = wb.create_sheet(title='品牌分類統計')
    ws_sum.append(['品牌名稱', '報價筆數', '供應商數', '最低報價(HKD)', '最高報價(HKD)', '平均報價(HKD)'])
    for brand in TARGET_BRANDS + ["其他品牌"]:
        subset = df_list[df_list['品牌分類'] == brand]
        if not subset.empty:
            ws_sum.append([brand, len(subset), subset['供應商'].nunique(),
                           int(subset['價格'].min()), int(subset['價格'].max()), round(subset['價格'].mean(), 1)])
    _style_header(ws_sum, 6, green_fill)
    for r in range(2, ws_sum.max_row + 1):
        for c in range(1, 7):
            cell = ws_sum.cell(row=r, column=c)
            cell.font, cell.border = (bold_font if c == 1 else body_font), thin_border
            if c >= 4:
                cell.number_format = '#,##0'
    ws_sum.auto_filter.ref = f"A1:F{ws_sum.max_row}"
    ws_sum.freeze_panes = 'A2'
    _autofit(ws_sum)

    wb.save(output)
    return output.getvalue()

# ==========================================
# 5. 篩選器狀態管理
# ==========================================
FILTER_KEYS = ['f_brands', 'f_vendors', 'f_price', 'f_kw', 'f_origin', 'f_stick', 'f_stock']

def reset_filters(price_bounds=None):
    """清除所有篩選條件（真正重置元件狀態）"""
    st.session_state['f_brands'] = []
    st.session_state['f_vendors'] = []
    st.session_state['f_kw'] = ""
    st.session_state['f_origin'] = []
    st.session_state['f_stick'] = '全部規格'
    st.session_state['f_stock'] = '全部（含詢價）'
    if price_bounds:
        st.session_state['f_price'] = price_bounds

def clear_results():
    for k in ['df_result', 'file_stats'] + FILTER_KEYS:
        st.session_state.pop(k, None)
    st.cache_data.clear()

# ==========================================
# 6. 版面：Hero 標題
# ==========================================
st.markdown(f"""
<div class="hero">
    <h1>🍂 雪茄批發報價自動彙整系統</h1>
    <p class="subtitle">
        上傳各家 Excel / PDF 報價單 → 自動辨識 30 個品牌與供應商 → 跨店比價 · 最低價排行 · 一鍵匯出 Excel
    </p>
    <span class="badge">⚡ {APP_VERSION} — 支援品牌維護 · 擷取支數 · 渠道標註</span>
</div>
""", unsafe_allow_html=True)

# ==========================================
# 7. 側邊欄：步驟 1 上傳與設定
# ==========================================
st.sidebar.markdown('<div class="sidebar-title">🍂 雪茄報價系統</div>', unsafe_allow_html=True)

# --- 品牌設定區 ---
with st.sidebar.expander("⚙️ 品牌關鍵字維護", expanded=False):
    st.caption("修改下方表格後點擊「儲存」即可立即套用。您也可以下載 CSV 修改後上傳至 GitHub 以永久保存。")
    edited_brands = st.data_editor(df_brand_rules, num_rows="dynamic", use_container_width=True, hide_index=True)
    if st.button("💾 儲存並套用新規則", use_container_width=True):
        edited_brands.to_csv(RULES_FILE, index=False, encoding='utf-8-sig')
        st.cache_data.clear()
        st.rerun()
    with open(RULES_FILE, "rb") as f:
        st.download_button("📥 下載規則檔", f, file_name=RULES_FILE, mime="text/csv", use_container_width=True)

st.sidebar.markdown('<div class="step-label">STEP 1 · 上傳報價單</div>', unsafe_allow_html=True)
uploaded_files = st.sidebar.file_uploader(
    "上傳報價單",
    type=["xlsx", "xls", "pdf"],
    accept_multiple_files=True,
    help="支援 Excel (.xlsx, .xls) 與 PDF，可一次上傳多個檔案",
    label_visibility="collapsed"
)

# ==========================================
# 8. 主畫面：空狀態
# ==========================================
if not uploaded_files and 'df_result' not in st.session_state:
    st.markdown("""
    <div class="empty-state">
        <div class="icon">📂</div>
        <div class="title">尚未上傳任何報價單</div>
        <div class="desc">👈 請在左側面板上傳 Excel 或 PDF 報價單，支援一次多檔上傳</div>
    </div>
    """, unsafe_allow_html=True)

    features = [
        ("🤖", "智慧解析引擎", "自動辨識 30 個品牌與供應商；並排報價、多工作表、PDF 皆可解析。"),
        ("📊", "跨店比價矩陣", "同一款雪茄各家報價橫向展開，最低價自動綠底標示。"),
        ("🏆", "最低價排行", "列出 2 家以上有報價的品項，依價差排序，找出最划算的進貨來源。"),
        ("📥", "多工作表 Excel", "清單、比價矩陣、最低價排行、品牌統計，一個檔案全部帶走。"),
    ]
    cols = st.columns(4)
    for col, (icon, title, desc) in zip(cols, features):
        col.markdown(f"""
        <div class="feature-card">
            <div class="f-icon">{icon}</div>
            <div class="f-title">{title}</div>
            <div class="f-desc">{desc}</div>
        </div>
        """, unsafe_allow_html=True)

    with st.expander("📖 使用說明", expanded=False):
        st.markdown("""
1. 在左側 **STEP 1** 上傳一或多份報價單（Excel / PDF）。
2. 在主畫面確認 **供應商名稱**，若自動辨識錯誤可直接點擊修改。
3. 按下 **🚀 開始解析**。
4. 使用左側 **STEP 3 篩選器** 過濾品牌、供應商、價格、產地或關鍵字。
5. 下載 Excel：會包含「全品項清單、跨店比價矩陣、最低價排行、品牌統計」4 個工作表。
        """)

# ==========================================
# 9. 主畫面：步驟 2 確認供應商 + 解析
# ==========================================
parse_button = False
if uploaded_files:
    st.markdown("""
    <div class="section-header">
        <span class="icon">🏷️</span>
        <span class="text">確認供應商名稱</span>
        <span class="desc">自動依檔名判斷 · 可點擊「供應商」欄位直接修改</span>
    </div>
    """, unsafe_allow_html=True)

    file_df = pd.DataFrame([{
        "檔名": f.name,
        "類型": "PDF" if f.name.lower().endswith('.pdf') else "Excel",
        "大小(KB)": round(f.size / 1024),
        "供應商": get_standard_vendor(f.name),
    } for f in uploaded_files])

    edited_files = st.data_editor(
        file_df,
        use_container_width=True,
        hide_index=True,
        disabled=["檔名", "類型", "大小(KB)"],
        column_config={"供應商": st.column_config.TextColumn("供應商 ✏️", required=True)},
        key=f"vendor_editor_{len(uploaded_files)}",
    )
    vendor_map = dict(zip(edited_files["檔名"], edited_files["供應商"]))

    st.sidebar.markdown('<div class="step-label">STEP 2 · 開始解析</div>', unsafe_allow_html=True)
    parse_button = st.sidebar.button("🚀 開始解析並合併轉檔", use_container_width=True, type="primary")
    st.sidebar.caption(f"📎 已選取 **{len(uploaded_files)}** 個檔案")

if parse_button and uploaded_files:
    all_records = []
    file_stats = []
    progress_bar = st.progress(0, text="⏳ 正在解析檔案...")

    for idx, file in enumerate(uploaded_files):
        vendor = (vendor_map.get(file.name) or get_standard_vendor(file.name)).strip()
        file_ext = file.name.rsplit('.', 1)[-1].lower()
        try:
            records = parse_file_v5(file.getvalue(), file_ext)
            for r in records:
                all_records.append({"供應商": vendor, **r})
            status = "✅ 成功" if records else "⚠️ 無資料"
            file_stats.append({"檔名": file.name, "供應商": vendor, "抓取筆數": len(records), "狀態": status})
        except Exception as e:
            file_stats.append({"檔名": file.name, "供應商": vendor, "抓取筆數": 0, "狀態": f"❌ {e}"})
        progress_bar.progress((idx + 1) / len(uploaded_files),
                              text=f"正在解析：{file.name} ({idx+1}/{len(uploaded_files)})")

    progress_bar.empty()

    # 新的解析結果 → 重置所有篩選條件，避免舊的價格範圍超出新資料
    for k in FILTER_KEYS:
        st.session_state.pop(k, None)
    st.session_state['file_stats'] = file_stats

    if all_records:
        df_result = pd.DataFrame(all_records)
        df_result['庫存盒數'] = pd.to_numeric(df_result['庫存盒數'], errors='coerce')
        brand_order_map = {b: i for i, b in enumerate(TARGET_BRANDS)}
        brand_order_map['其他品牌'] = 999
        df_result['排序權重'] = df_result['品牌分類'].map(brand_order_map)
        # 依品牌 → 品名 → 價格排序，去重時保留同供應商同品名的最低價
        df_result = df_result.sort_values(by=['排序權重', '品名規格', '價格'])
        df_result = df_result.drop_duplicates(subset=['供應商', '品牌分類', '品名規格']).drop(columns=['排序權重'])
        st.session_state['df_result'] = df_result.reset_index(drop=True)
        st.toast(f"✅ 解析完成，共 {len(df_result)} 筆資料", icon="🎉")
    else:
        st.session_state.pop('df_result', None)
        st.warning("⚠️ 所有檔案均未成功解析出有效資料，請檢查檔案格式。")

# ==========================================
# 10. 結果展示區
# ==========================================
if 'df_result' in st.session_state:
    df_result = st.session_state['df_result']

    # --- 側邊欄：步驟 3 篩選器 ---
    st.sidebar.markdown('<div class="step-label">STEP 3 · 篩選資料</div>', unsafe_allow_html=True)

    price_min = int(df_result['價格'].min())
    price_max = int(df_result['價格'].max())
    if price_max <= price_min:          # 所有價格相同時，滑桿需要不同的上下限
        price_max = price_min + 50
    price_bounds = (price_min, price_max)

    # 首次顯示時初始化篩選元件狀態
    if 'f_brands' not in st.session_state:
        reset_filters(price_bounds)

    brand_order = [b for b in TARGET_BRANDS + ['其他品牌'] if b in set(df_result['品牌分類'])]
    st.sidebar.multiselect("📌 品牌分類", brand_order, key='f_brands', placeholder="全部品牌")
    st.sidebar.multiselect("🏬 供應商", sorted(df_result['供應商'].unique()), key='f_vendors', placeholder="全部供應商")
    
    # 依照您提供的對照字典順序來排列渠道選項
    standard_channels = list(dict.fromkeys(CHANNEL_CLEAN_MAP.values()))
    available_channels = [c for c in standard_channels if c in set(df_result['渠道/產地'])]
    other_channels = sorted([c for c in df_result['渠道/產地'].unique() if c and c not in available_channels])
    origins = available_channels + other_channels
    
    if origins:
        st.sidebar.multiselect("🌍 渠道/產地", origins, key='f_origin', placeholder="全部渠道/產地")
        
    stick_options = [
        '全部規格', '25支常規盒', '10支精裝', '12支(千里達)', '15支鋁管/盒', 
        '20支裝', '50支滑蓋櫃', '機制條裝(60/100支)', '其他/未明確標註'
    ]
    st.sidebar.selectbox('📦 包裝支數規格', stick_options, key='f_stick')

    st.sidebar.radio(
        '📦 現貨狀態', ['全部（含詢價）', '僅看在席現貨 (庫存 > 0)', '大宗現貨 (≥ 5盒)'], key='f_stock'
    )

    st.sidebar.slider("💰 價格區間 (HKD)", min_value=price_min, max_value=price_max,
                      step=50, format="$%d", key='f_price')
    st.sidebar.text_input("🔎 品名搜尋", key='f_kw', placeholder="短丘、D4、BHK... (空格分隔可多關鍵字)")

    c1, c2 = st.sidebar.columns(2)
    c1.button("🗑️ 清除篩選", use_container_width=True, on_click=reset_filters, args=(price_bounds,))
    c2.button("♻️ 重新開始", use_container_width=True, on_click=clear_results)

    # --- 套用篩選 ---
    selected_brands = st.session_state['f_brands']
    selected_vendors = st.session_state['f_vendors']
    selected_origins = st.session_state.get('f_origin', [])
    selected_stick = st.session_state.get('f_stick', '全部規格')
    selected_stock = st.session_state.get('f_stock', '全部（含詢價）')
    price_range = st.session_state['f_price']
    keyword = st.session_state['f_kw'].strip()

    df_display = df_result
    if selected_brands:
        df_display = df_display[df_display['品牌分類'].isin(selected_brands)]
    if selected_vendors:
        df_display = df_display[df_display['供應商'].isin(selected_vendors)]
    if selected_origins:
        df_display = df_display[df_display['渠道/產地'].isin(selected_origins)]
        
    if selected_stick == '25支常規盒': df_display = df_display[df_display['規格支數'] == 25]
    elif selected_stick == '10支精裝': df_display = df_display[df_display['規格支數'] == 10]
    elif selected_stick == '12支(千里達)': df_display = df_display[df_display['規格支數'] == 12]
    elif selected_stick == '15支鋁管/盒': df_display = df_display[df_display['規格支數'] == 15]
    elif selected_stick == '20支裝': df_display = df_display[df_display['規格支數'] == 20]
    elif selected_stick == '50支滑蓋櫃': df_display = df_display[df_display['規格支數'] == 50]
    elif selected_stick == '機制條裝(60/100支)': df_display = df_display[df_display['規格支數'].isin([60, 100])]
    elif selected_stick == '其他/未明確標註': df_display = df_display[df_display['規格支數'].isna()]

    if selected_stock == '僅看在席現貨 (庫存 > 0)': df_display = df_display[df_display['庫存盒數'] > 0]
    elif selected_stock == '大宗現貨 (≥ 5盒)': df_display = df_display[df_display['庫存盒數'] >= 5]
        
    df_display = df_display[(df_display['價格'] >= price_range[0]) & (df_display['價格'] <= price_range[1])]
    if keyword:
        # 多關鍵字：以空格分隔，需同時符合（AND）
        for kw in keyword.split():
            df_display = df_display[df_display['品名規格'].str.contains(kw, case=False, na=False, regex=False)]
    df_display = df_display.reset_index(drop=True)

    # --- 狀態列（顯示目前的篩選條件） ---
    chips = []
    if selected_brands: chips.append(f"品牌：{'、'.join(selected_brands)}")
    if selected_vendors: chips.append(f"供應商：{'、'.join(selected_vendors)}")
    if selected_origins: chips.append(f"渠道：{'、'.join(selected_origins)}")
    if tuple(price_range) != price_bounds: chips.append(f"價格：${price_range[0]:,} ~ ${price_range[1]:,}")
    if keyword: chips.append(f"搜尋：{keyword}")
    chips_html = "".join(f'<span class="chip">{c}</span>' for c in chips)
    st.markdown(f"""
    <div class="status-bar">
        <div class="dot"></div>
        <span class="text">已載入 {len(df_result):,} 筆 → 顯示 {len(df_display):,} 筆</span>
        {chips_html}
    </div>
    """, unsafe_allow_html=True)

    # --- 解析報告 ---
    if 'file_stats' in st.session_state:
        stats_df = pd.DataFrame(st.session_state['file_stats'])
        failed = stats_df[~stats_df['狀態'].str.startswith('✅')]
        with st.expander(f"📋 解析報告：{len(stats_df)} 個檔案" + (f"（⚠️ {len(failed)} 個需檢查）" if len(failed) else ""),
                         expanded=bool(len(failed))):
            st.dataframe(stats_df, use_container_width=True, hide_index=True,
                         column_config={"抓取筆數": st.column_config.ProgressColumn(
                             "抓取筆數", format="%d", min_value=0,
                             max_value=int(max(stats_df['抓取筆數'].max(), 1)))})
            st.caption(f"原始抓取 {stats_df['抓取筆數'].sum():,} 筆 → 去除同供應商重複品項後 {len(df_result):,} 筆")

    # --- KPI ---
    st.markdown("""
    <div class="section-header">
        <span class="icon">📈</span>
        <span class="text">數據總覽</span>
        <span class="desc">隨篩選條件即時更新</span>
    </div>
    """, unsafe_allow_html=True)

    df_best = build_best_price(df_display)
    k1, k2, k3, k4, k5, k6 = st.columns(6)
    k1.metric("篩選品項", f"{len(df_display):,} 筆")
    k2.metric("供應商", f"{df_display['供應商'].nunique()} 家")
    k3.metric("品牌", f"{df_display['品牌分類'].nunique()} 個")
    if not df_display.empty:
        k4.metric("最低價", f"${int(df_display['價格'].min()):,}")
        k5.metric("平均價", f"${int(df_display['價格'].mean()):,}")
        k6.metric("可比價品項", f"{len(df_best):,} 款", help="同一品名在 2 家以上供應商有報價的款數")
    else:
        k4.metric("最低價", "—"); k5.metric("平均價", "—"); k6.metric("可比價品項", "—")

    # --- 下載區 ---
    today_str = datetime.now().strftime("%Y%m%d")
    pivot_df = build_pivot(df_display)
    dl1, dl2, dl3 = st.columns([3, 2, 3])
    with dl1:
        if not df_display.empty:
            with st.spinner("正在產生 Excel..."):
                excel_bytes = generate_excel(df_display, df_best, pivot_df)
            st.download_button(
                label="📥 下載完整 Excel（4 個工作表）",
                data=excel_bytes,
                file_name=f"雪茄報價彙總_{today_str}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
    with dl2:
        if not df_display.empty:
            st.download_button(
                label="📄 下載 CSV",
                data=df_display.to_csv(index=False).encode('utf-8-sig'),
                file_name=f"雪茄報價彙總_{today_str}.csv",
                mime="text/csv",
                use_container_width=True,
            )
    with dl3:
        st.caption("💡 下載內容會套用目前的篩選條件。Excel 含：全品項清單、跨店比價矩陣、最低價排行、品牌統計。")

    # --- 資料檢視分頁 ---
    st.markdown("""
    <div class="section-header">
        <span class="icon">📋</span>
        <span class="text">資料檢視</span>
        <span class="desc">切換標籤查看不同視圖</span>
    </div>
    """, unsafe_allow_html=True)

    other_count = int((df_display['品牌分類'] == '其他品牌').sum())
    tab_list, tab_matrix, tab_best, tab_chart, tab_other = st.tabs([
        f"📋 全品項清單 ({len(df_display)})",
        "📊 跨店比價矩陣",
        f"🏆 最低價排行 ({len(df_best)})",
        "📈 統計圖表",
        f"⚠️ 未分類品項 ({other_count})",
    ])

    no_data_msg = "⚠️ 目前篩選條件下無資料，請調整左側篩選器。"

    with tab_matrix:
        if pivot_df.empty:
            st.warning(no_data_msg)
        else:
            only_multi = st.toggle("只顯示 2 家以上有報價的品項", value=False, key="matrix_only_multi")
            show_pivot = pivot_df[pivot_df.notna().sum(axis=1) >= 2] if only_multi else pivot_df
            st.caption(f"💡 綠底粗體 = 該品項的最低報價 · 共 {len(show_pivot):,} 款")

            def highlight_min(row):
                valid = row.dropna()
                if len(valid) < 2:
                    return ['' for _ in row]
                min_val = valid.min()
                return ['background-color: #D4EDDA; color: #14532D; font-weight: bold'
                        if pd.notna(v) and v == min_val else '' for v in row]

            if show_pivot.size <= 200_000:
                st.dataframe(show_pivot.style.apply(highlight_min, axis=1).format("{:,}", na_rep=""),
                             use_container_width=True, height=600)
            else:
                st.dataframe(show_pivot, use_container_width=True, height=600)

    with tab_best:
        if df_best.empty:
            st.info("目前篩選條件下，沒有同一品名在 2 家以上供應商同時有報價的品項。")
        else:
            st.caption("同一品名在 2 家以上供應商有報價時，列出最低價來源與價差（依價差由大到小排序）。")
            st.dataframe(
                df_best, use_container_width=True, height=600, hide_index=True,
                column_config={
                    "最低價": st.column_config.NumberColumn(format="$%d"),
                    "最高價": st.column_config.NumberColumn(format="$%d"),
                    "價差": st.column_config.NumberColumn(format="$%d"),
                    "價差%": st.column_config.ProgressColumn(
                        format="%.1f%%", min_value=0,
                        max_value=float(max(df_best['價差%'].max(), 1))),
                },
            )

    with tab_list:
        if df_display.empty:
            st.warning(no_data_msg)
        else:
            st.dataframe(
                df_display, use_container_width=True, height=600, hide_index=True,
                column_order=['供應商', '品牌分類', '品名規格', '規格支數', '庫存盒數', '價格', '渠道/產地'],
                column_config={
                    "價格": st.column_config.NumberColumn("整盒批發價", format="$%d"),
                    "規格支數": st.column_config.NumberColumn("規格支數", format="%d 支"),
                    "庫存盒數": st.column_config.NumberColumn("庫存現貨", format="%d 盒"),
                    "品名規格": st.column_config.TextColumn(width="large"),
                },
            )

    with tab_chart:
        if df_display.empty:
            st.warning(no_data_msg)
        else:
            ch1, ch2 = st.columns(2)
            with ch1:
                st.markdown("**各品牌報價筆數**")
                st.bar_chart(df_display['品牌分類'].value_counts(), horizontal=True, color="#2C5364")
            with ch2:
                st.markdown("**各供應商報價筆數**")
                st.bar_chart(df_display['供應商'].value_counts(), horizontal=True, color="#2C5E3B")
            st.markdown("**各品牌平均報價 (HKD)**")
            st.bar_chart(df_display.groupby('品牌分類')['價格'].mean().round(0).sort_values(ascending=False),
                         color="#1F4E78")

    with tab_other:
        others = df_display[df_display['品牌分類'] == '其他品牌']
        if others.empty:
            st.success("🎉 目前所有品項都已成功分類到 30 個品牌中！")
        else:
            st.caption("以下品項沒有匹配到 30 個品牌的關鍵字。可能是非古巴品牌、雜項（如雪茄剪、保濕盒），"
                       "或需要補充品牌關鍵字。把常見的品名告訴我，我可以幫你加入分類規則。")
            st.dataframe(others, use_container_width=True, height=500, hide_index=True,
                         column_config={"價格": st.column_config.NumberColumn("價格 (HKD)", format="$%d")})

# --- 頁腳 ---
st.sidebar.markdown("---")
st.sidebar.caption(f"🍂 {APP_VERSION} · {datetime.now().strftime('%Y-%m-%d %H:%M')}")
st.markdown(f"""
<div class="footer">
    🍂 雪茄批發報價自動彙整系統 {APP_VERSION} · Built with Streamlit
</div>
""", unsafe_allow_html=True)
